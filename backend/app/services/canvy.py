import io
import json
import logging
import secrets
from pathlib import Path

from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Canva, SzablonCanvy
from app.repositories.kreator import CanvaRepository, FiszkaRepository
from app.schemas.kreator import CanvaRead, CanvaUpdate
from app.schemas.kreator import SzablonCanvy as SzablonRead
from app.services.errors import KreatorError

logger = logging.getLogger(__name__)

TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "data"
MAX_VALUE = 3000


async def import_templates(session: AsyncSession, directory: Path = TEMPLATES_DIR) -> int:
    """Szablony canv to dane (pliki JSON w `app/data`): dopisuje brakujące, nie nadpisuje zmian."""
    repo = CanvaRepository(session)
    added = 0
    for file in sorted(directory.glob("canva_*.json")):
        data = json.loads(file.read_text(encoding="utf-8"))
        if await repo.template(data["slug"]) is None:
            session.add(SzablonCanvy(**data))
            added += 1
    if added:
        await session.commit()
        logger.info("Dodano %s szablonów canvy", added)
    return added


class CanvaService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.repo = CanvaRepository(session)
        self.fiszki = FiszkaRepository(session)

    async def templates(self) -> list[SzablonRead]:
        rows = await self.repo.templates()
        return [SzablonRead.model_validate(t, from_attributes=True) for t in rows]

    async def create(self, slug: str, title: str, fiszka_token: str | None) -> Canva:
        template = await self.repo.template(slug)
        if template is None:
            raise KreatorError("Nie znaleziono takiej canvy.", 404)
        fiszka_id = None
        if fiszka_token:
            fiszka = await self.fiszki.get(fiszka_token)
            if fiszka is None:
                raise KreatorError("Nie znaleziono fiszki do powiązania.", 404)
            fiszka_id = fiszka.id
        canva = Canva(
            token=secrets.token_urlsafe(24),
            szablon_slug=slug,
            tytul=title.strip(),
            fiszka_id=fiszka_id,
            wartosci={},
        )
        await self.repo.add(canva)
        await self.session.commit()
        await self.session.refresh(canva)
        return canva

    async def get(self, token: str) -> Canva:
        canva = await self.repo.get(token)
        if canva is None:
            raise KreatorError("Nie znaleziono canvy. Sprawdź adres albo zacznij nową.", 404)
        return canva

    async def update(self, token: str, data: CanvaUpdate) -> Canva:
        canva = await self.get(token)
        if data.tytul is not None:
            canva.tytul = data.tytul.strip()
        if data.wartosci is not None:
            template = await self.repo.template(canva.szablon_slug)
            keys = {s["klucz"] for s in template.sekcje} if template else set()
            values = dict(canva.wartosci)
            for key, text in data.wartosci.items():
                if key not in keys:
                    raise KreatorError("Nieznane pole canvy.", 422)
                if len(text) > MAX_VALUE:
                    raise KreatorError(f"Tekst w polu przekracza {MAX_VALUE} znaków.", 422)
                values[key] = text
            canva.wartosci = values
        await self.session.commit()
        await self.session.refresh(canva)
        return canva

    async def read(self, canva: Canva) -> CanvaRead:
        template = await self.repo.template(canva.szablon_slug)
        assert template is not None
        fiszka = await self.fiszki.by_id(canva.fiszka_id) if canva.fiszka_id else None
        return CanvaRead(
            token=canva.token,
            tytul=canva.tytul,
            szablon=SzablonRead.model_validate(template, from_attributes=True),
            wartosci=canva.wartosci or {},
            fiszka_token=fiszka.token if fiszka else None,
            syntetyczna=canva.syntetyczna,
            updated_at=canva.updated_at,
        )

    async def export_docx(self, canva: Canva) -> bytes:
        read = await self.read(canva)
        doc = Document()
        doc.add_heading(read.tytul or read.szablon.nazwa, level=1)
        group = None
        for section in read.szablon.sekcje:
            if section.grupa != group:
                group = section.grupa
                doc.add_heading(group, level=2)
            doc.add_heading(section.tytul, level=3)
            doc.add_paragraph(read.wartosci.get(section.klucz, "").strip() or "[DO UZUPEŁNIENIA]")
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

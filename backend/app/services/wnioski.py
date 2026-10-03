import io
import secrets
from datetime import date

from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Nabor, Wniosek
from app.repositories.kreator import FiszkaRepository, NaborRepository, WniosekRepository
from app.schemas.kreator import (
    POLA_FISZKI_ETYKIETY,
    WniosekPoleRead,
    WniosekRead,
)
from app.schemas.ticket import TicketCreate
from app.services.ai import AIGateway
from app.services.errors import KreatorError
from app.services.fiszki import FiszkaService
from app.services.kreator_ai import draft_wniosek, values_of
from app.services.nabory import NaborService, status_of, to_read
from app.services.tickets import TicketService

PLACEHOLDER = "[DO UZUPEŁNIENIA]"
TICKET_LIMIT = 20000


def _field_source(stored: dict, text: str) -> str:
    """„ai” i „fiszka” tylko dopóki tekst jest taki, jak wygenerowany; zmiana = praca autora."""
    if stored["zrodlo"] == "brak" and not text:
        return "brak"
    if text != stored.get("wygenerowany", ""):
        return "uzytkownik"
    return stored["zrodlo"]


class WniosekService:
    def __init__(
        self,
        session: AsyncSession,
        ai: AIGateway,
        fiszki: FiszkaService,
        nabory: NaborService,
        tickets: TicketService,
    ) -> None:
        self.session = session
        self.ai = ai
        self.fiszki = fiszki
        self.nabory = nabory
        self.tickets = tickets
        self.repo = WniosekRepository(session)
        self.fiszki_repo = FiszkaRepository(session)
        self.nabory_repo = NaborRepository(session)

    async def create(self, fiszka_token: str, nabor_slug: str, today: date) -> Wniosek:
        """Jeden wniosek na parę (fiszka, nabór): powrót daje ten sam szkic, bez gubienia zmian."""
        fiszka = await self.fiszki.get(fiszka_token)
        nabor = await self.nabory.get_active(nabor_slug, today)
        existing = await self.repo.for_pair(fiszka.id, nabor.id)
        if existing is not None:
            return existing
        values = values_of(fiszka, self.fiszki.categories())
        if not values:
            raise KreatorError("Najpierw opisz pomysł w fiszce, wtedy przygotujemy wniosek.", 422)
        pola, message = await draft_wniosek(self.ai, nabor, values)
        wniosek = Wniosek(
            token=secrets.token_urlsafe(24),
            fiszka_id=fiszka.id,
            nabor_id=nabor.id,
            pola=pola,
            status="szkic",
            komunikat_ai=message,
        )
        await self.repo.add(wniosek)
        await self.session.commit()
        await self.session.refresh(wniosek)
        return wniosek

    async def get(self, token: str) -> Wniosek:
        wniosek = await self.repo.get(token)
        if wniosek is None:
            raise KreatorError("Nie znaleziono wniosku. Sprawdź adres.", 404)
        return wniosek

    async def update(self, token: str, texts: dict[str, str]) -> Wniosek:
        wniosek = await self.get(token)
        if wniosek.status == "wyslany":
            raise KreatorError("Ten wniosek został już wysłany i nie można go zmieniać.")
        nabor = await self.nabory_repo.by_id(wniosek.nabor_id)
        limits = {p["klucz"]: int(p["limit"]) for p in nabor.pola} if nabor else {}
        pola = {k: dict(v) for k, v in wniosek.pola.items()}
        for key, text in texts.items():
            if key not in pola:
                raise KreatorError("Nieznane pole wniosku.", 422)
            if len(text) > limits.get(key, 10**6):
                raise KreatorError(f"Tekst w polu przekracza limit {limits[key]} znaków.", 422)
            pola[key]["tekst"] = text
        wniosek.pola = pola
        await self.session.commit()
        await self.session.refresh(wniosek)
        return wniosek

    async def read(self, wniosek: Wniosek, today: date) -> WniosekRead:
        nabor = await self.nabory_repo.by_id(wniosek.nabor_id)
        fiszka = await self.fiszki_repo.by_id(wniosek.fiszka_id)
        assert nabor is not None and fiszka is not None
        return WniosekRead(
            token=wniosek.token,
            fiszka_token=fiszka.token,
            nabor=to_read(nabor, today),
            nabor_aktywny=status_of(nabor, today) == "aktywny",
            pola=self._fields(nabor, wniosek),
            kryteria=nabor.kryteria,
            status=wniosek.status,  # type: ignore[arg-type]
            token_watku=wniosek.token_watku,
            komunikat_ai=wniosek.komunikat_ai,
            updated_at=wniosek.updated_at,
        )

    def _fields(self, nabor: Nabor, wniosek: Wniosek) -> list[WniosekPoleRead]:
        fields = []
        for pole in nabor.pola:
            stored = wniosek.pola.get(pole["klucz"]) or {
                "tekst": "",
                "wygenerowany": "",
                "zrodlo": "brak",
                "uzyte_pola": [],
            }
            text = stored["tekst"]
            fields.append(
                WniosekPoleRead(
                    klucz=pole["klucz"],
                    etykieta=pole["etykieta"],
                    limit=int(pole["limit"]),
                    wskazowka=pole.get("wskazowka", ""),
                    tekst=text,
                    zrodlo=_field_source(stored, text),  # type: ignore[arg-type]
                    uzyte_pola=[POLA_FISZKI_ETYKIETY[s] for s in stored.get("uzyte_pola", [])],
                    do_uzupelnienia=not text.strip(),
                )
            )
        return fields

    async def send(self, token: str, today: date) -> str:
        """Opcjonalne wysłanie do ROPS jako zwykłe zgłoszenie w skrzynce panelu."""
        wniosek = await self.get(token)
        if wniosek.status == "wyslany" and wniosek.token_watku:
            return wniosek.token_watku
        nabor = await self.nabory_repo.by_id(wniosek.nabor_id)
        assert nabor is not None
        if status_of(nabor, today) != "aktywny":
            raise KreatorError(self.nabory.inactive_message(nabor, today))
        fiszka = await self.fiszki_repo.by_id(wniosek.fiszka_id)
        lines = [f"[Wniosek z Kreatora] Nabór: {nabor.nazwa}", ""]
        for pole in self._fields(nabor, wniosek):
            lines += [f"{pole.etykieta}:", pole.tekst.strip() or PLACEHOLDER, ""]
        # Wniosek bywa dłuższy niż limit publicznego formularza, więc bez walidacji długości.
        data = TicketCreate.model_construct(
            tresc="\n".join(lines)[:TICKET_LIMIT], autor_nazwa=None, autor_email=None
        )
        ticket = await self.tickets.create(data, synthetic=bool(fiszka and fiszka.syntetyczna))
        wniosek.token_watku = ticket.token_watku
        wniosek.status = "wyslany"
        await self.session.commit()
        return ticket.token_watku

    async def export_text(self, wniosek: Wniosek, today: date) -> tuple[str, str]:
        read = await self.read(wniosek, today)
        lines = [read.nabor.nazwa, "Szkic wniosku przygotowany w Kreatorze pomysłów (Splot)", ""]
        for pole in read.pola:
            lines += [pole.etykieta, pole.tekst.strip() or PLACEHOLDER, ""]
        return read.nabor.slug, "\n".join(lines)

    async def export_docx(self, wniosek: Wniosek, today: date) -> tuple[str, bytes]:
        read = await self.read(wniosek, today)
        doc = Document()
        doc.add_heading(read.nabor.nazwa, level=1)
        doc.add_paragraph(
            "Szkic wniosku przygotowany w Kreatorze pomysłów (Splot). "
            f"Pola oznaczone {PLACEHOLDER} wymagają uzupełnienia."
        )
        for pole in read.pola:
            doc.add_heading(pole.etykieta, level=2)
            doc.add_paragraph(pole.tekst.strip() or PLACEHOLDER)
        buffer = io.BytesIO()
        doc.save(buffer)
        return read.nabor.slug, buffer.getvalue()


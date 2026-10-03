"""Generuje syntetyczne dokumenty projektu do demo uploadu (assets/demo/*.docx).

Uruchomienie z katalogu backend: python scripts/make_demo_documents.py
Treść jest w całości wymyślona (żadnych prawdziwych organizacji ani osób).
"""

from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parents[2] / "assets" / "demo"

DOCUMENTS = {
    "dokument-projektu-1-autobus-zdrowia.docx": [
        ("title", "Sąsiedzki Autobus Zdrowia"),
        ("note", "DOKUMENT SYNTETYCZNY przygotowany na potrzeby demonstracji. Dane są wymyślone."),
        ("h", "1. Problem"),
        (
            "p",
            "Seniorzy mieszkający w małych wsiach nie mają jak dojechać do przychodni i do "
            "apteki, ponieważ w ich miejscowościach nie kursuje komunikacja publiczna.",
        ),
        ("h", "2. Odbiorcy"),
        ("p", "Osoby w wieku 65 lat i więcej z gmin wiejskich bez regularnej komunikacji."),
        ("h", "3. Opis rozwiązania"),
        (
            "p",
            "Gmina uruchamia mały autobus, który dwa razy w tygodniu zbiera mieszkańców z "
            "ustalonych przystanków i wozi ich do przychodni i apteki. Kursy prowadzą "
            "wolontariusze z lokalnych kół gospodyń wiejskich.",
        ),
        ("h", "4. Wyniki pilotażu"),
        (
            "p",
            "Pilotaż w dwóch gminach trwał sześć miesięcy i objął 62 seniorów. Liczba "
            "odwołanych wizyt lekarskich spadła o 40 procent.",
        ),
        ("h", "5. Koszt i wdrożenie"),
        ("p", "Koszt wdrożenia dla gminy jest niski, ponieważ wykorzystuje istniejący pojazd."),
        ("p", "Uruchomienie zajmuje kilka tygodni."),
        (
            "p",
            "Wymagania: pojazd gminny, jeden koordynator z urzędu gminy oraz grupa wolontariuszy.",
        ),
        ("p", "Autor: Fundacja Wiejska Mobilność (organizacja wymyślona)."),
    ],
    "dokument-projektu-2-telefon-opiekunow.docx": [
        ("title", "Telefon Wsparcia dla Opiekunów"),
        ("note", "DOKUMENT SYNTETYCZNY przygotowany na potrzeby demonstracji. Dane są wymyślone."),
        ("h", "1. Problem"),
        (
            "p",
            "Opiekunowie osób z demencją są przeciążeni i nie mają gdzie szybko zapytać o radę, "
            "co prowadzi do wypalenia i przedwczesnego umieszczania bliskich w placówkach.",
        ),
        ("h", "2. Odbiorcy"),
        ("p", "Rodziny i opiekunowie nieformalni osób z demencją."),
        ("h", "3. Opis rozwiązania"),
        (
            "p",
            "Dyżurny telefon, pod którym przeszkoleni pracownicy ośrodka pomocy społecznej "
            "udzielają porad i kierują do lokalnych usług. Rozwiązanie jest na etapie pomysłu, "
            "nie było jeszcze testowane.",
        ),
        ("h", "4. Wdrożenie"),
        ("p", "Wymagania: dwie przeszkolone osoby i numer telefonu. Koszt nie został oszacowany."),
    ],
}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for filename, blocks in DOCUMENTS.items():
        doc = Document()
        for kind, text in blocks:
            if kind == "title":
                doc.add_heading(text, level=0)
            elif kind == "h":
                doc.add_heading(text, level=2)
            else:
                doc.add_paragraph(text)
        doc.save(OUT / filename)
        print("zapisano", OUT / filename)


if __name__ == "__main__":
    main()
